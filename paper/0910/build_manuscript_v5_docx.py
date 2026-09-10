"""Build the reviewable Word version of manuscript_v5.md."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "paper/0910/manuscript_v5.md"
OUTPUT = ROOT / "paper/0910/manuscript_v5.docx"

INLINE = re.compile(
    r"(\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`|\*([^*]+)\*)"
)

PLACEHOLDERS = {
    "TandemX addresses a distinction": "[INSERT FIGURE 1 NEAR HERE]",
    "We evaluated repeat discovery": "[INSERT FIGURE 2 NEAR HERE]",
    "We next measured how discovery": "[INSERT SUPPLEMENTARY FIGURE S1 NEAR HERE]",
    "We next asked whether the workflow": "[INSERT FIGURE 3 NEAR HERE]",
    "The first biological test asked": "[INSERT FIGURE 4 NEAR HERE]",
    "We selected six newer-assembly": "[INSERT FIGURE 5 NEAR HERE]",
}


def set_font(run, name: str = "Times New Roman", size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    properties.append(fonts)
    run.append(properties)
    content = OxmlElement("w:t")
    content.text = text
    run.append(content)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE.finditer(text):
        if match.start() > cursor:
            set_font(paragraph.add_run(text[cursor : match.start()]))
        if match.group(2) is not None:
            add_hyperlink(paragraph, match.group(2), match.group(3))
        elif match.group(4) is not None:
            run = paragraph.add_run(match.group(4))
            set_font(run)
            run.bold = True
        elif match.group(5) is not None:
            run = paragraph.add_run(match.group(5))
            set_font(run)
        else:
            run = paragraph.add_run(match.group(6))
            set_font(run)
            run.italic = True
        cursor = match.end()
    if cursor < len(text):
        set_font(paragraph.add_run(text[cursor:]))


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 90, bottom: int = 80, end: int = 90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        borders.append(item)


def add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = {
        5: [1.15, 1.10, 1.35, 1.40, 1.45],
        7: [0.95, 0.75, 1.35, 0.90, 0.80, 0.85, 0.85],
    }[len(rows[0])]
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.width = Inches(widths[column_index])
            cell._tc.get_or_add_tcPr().tcW.set(
                qn("w:w"), str(round(widths[column_index] * 1440))
            )
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(
                cell,
                top=45 if len(rows[0]) == 7 else 80,
                bottom=45 if len(rows[0]) == 7 else 80,
            )
            if row_index == 0:
                set_cell_shading(cell, "17365D")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "EDF3F8")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if column_index > 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, value)
            for run in paragraph.runs:
                set_font(run, size=7.5 if len(rows[0]) == 7 else 8.2)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_styles(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(0.87)
    section.bottom_margin = Inches(0.87)
    section.left_margin = Inches(0.91)
    section.right_margin = Inches(0.91)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    settings = {
        "Title": (18, True, 1.0, 12, 6),
        "Heading 1": (14, True, 1.0, 12, 6),
        "Heading 2": (12, True, 1.0, 12, 6),
    }
    for name, (size, bold, spacing, before, after) in settings.items():
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = spacing
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        borders = style._element.get_or_add_pPr().find(qn("w:pBdr"))
        if borders is not None:
            style._element.get_or_add_pPr().remove(borders)


def add_placeholder(document: Document, label: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(label)
    set_font(run, size=10.5)
    run.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    index = start
    while index < len(lines) and lines[index].startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_styles(document)
    index = 0
    paragraph_buffer: list[str] = []

    def flush() -> None:
        if not paragraph_buffer:
            return
        text = " ".join(part.strip() for part in paragraph_buffer)
        paragraph_buffer.clear()
        paragraph = document.add_paragraph()
        add_inline(paragraph, text)
        if text.startswith("**Table "):
            paragraph.paragraph_format.keep_with_next = True
        for prefix, label in PLACEHOLDERS.items():
            if text.startswith(prefix):
                add_placeholder(document, label)
                break

    while index < len(lines):
        line = lines[index]
        if not line.strip():
            flush()
            index += 1
            continue
        if line.startswith("|"):
            flush()
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        if line.startswith("# "):
            flush()
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            borders = paragraph._p.get_or_add_pPr().find(qn("w:pBdr"))
            if borders is not None:
                paragraph._p.get_or_add_pPr().remove(borders)
            add_inline(paragraph, line[2:].strip())
            index += 1
            continue
        if line.startswith("## "):
            flush()
            heading = line[3:].strip()
            if heading == "Figure legends":
                document.add_page_break()
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, heading)
            index += 1
            continue
        if line.startswith("### "):
            flush()
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline(paragraph, line[4:].strip())
            index += 1
            continue
        if re.match(r"^\d+\. ", line):
            flush()
        paragraph_buffer.append(line)
        index += 1
    flush()
    document.core_properties.title = lines[0].removeprefix("# ")
    document.core_properties.subject = "TandemX manuscript version 5"
    document.core_properties.keywords = "TandemX, satellite DNA, genome assembly"
    document.save(OUTPUT)


if __name__ == "__main__":
    build()
